"""docx_review.py: each check bites on a planted defect and stays quiet on a clean file.

Gates prove they bite (CLAUDE.md): every rule below is seeded into an otherwise clean document,
and the clean document itself must pass, so a rule that never fires cannot read as green.
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

docx = pytest.importorskip("docx", reason="python-docx is in the dev extra")
from docx.oxml import OxmlElement, parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from docx.shared import Mm  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import docx_review as dr  # noqa: E402

OWNER = "Example Pty Ltd"


def _page_field(par) -> None:
    for kind in ("begin", "instr", "end"):
        r = OxmlElement("w:r")
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.text = " PAGE "
            r.append(it)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r.append(fc)
        par._p.append(r)


def clean_doc():
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    _page_field(s.footer.paragraphs[0])
    d.core_properties.title = "A clean document"
    d.core_properties.author = OWNER
    d.core_properties.last_modified_by = OWNER
    d.add_heading("Heading", 1)
    d.add_paragraph("Plain body text with a single space between words.")
    d.add_paragraph("First item", style="List Bullet")
    t = d.add_table(rows=3, cols=2)
    t.style = d.styles["Table Grid"]
    for i, row in enumerate(t.rows):
        row.cells[0].text, row.cells[1].text = f"key {i}", f"value {i}"
    return d


def save(d, tmp_path: Path, name: str = "doc.docx") -> Path:
    p = tmp_path / name
    d.save(str(p))
    return p


def rules(path: Path, **kw) -> set[str]:
    return {f.rule for f in dr.check(path, owner=OWNER, **kw).findings}


def test_clean_document_passes(tmp_path):
    rep = dr.check(save(clean_doc(), tmp_path), owner=OWNER)
    assert rep.status == "pass", [f.rule for f in rep.findings]
    assert rep.checks_run >= 12  # a denominator: zero checks must never read as a pass


def test_page_size_missing_and_wrong(tmp_path):
    d = clean_doc()
    sectPr = d.sections[0]._sectPr
    sectPr.remove(sectPr.find(qn("w:pgSz")))
    assert "page-size-missing" in rules(save(d, tmp_path))
    d = clean_doc()
    d.sections[0].page_width, d.sections[0].page_height = Mm(215.9), Mm(279.4)
    p = save(d, tmp_path, "letter.docx")
    assert "page-size" in rules(p)
    assert "page-size" not in rules(p, page="letter")


def test_no_page_numbers(tmp_path):
    d = clean_doc()
    f = d.sections[0].footer.paragraphs[0]._p
    for r in list(f):
        if r.tag == qn("w:r"):
            f.remove(r)
    assert "no-page-numbers" in rules(save(d, tmp_path))


def test_fragile_font(tmp_path):
    d = clean_doc()
    d.styles["Normal"].font.name = "Aptos"
    assert "fragile-font" in rules(save(d, tmp_path))


BOX = (
    f'<w:r {nsdecls("w", "wp", "a")}><w:drawing><wp:inline><wp:extent cx="6276975" cy="314325"/>'
    '<wp:docPr id="9" name="Box"/>'
    '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
    "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>"
)


def test_empty_shapes_in_body_and_footer(tmp_path):
    d = clean_doc()
    d.add_paragraph()._p.append(parse_xml(BOX))
    d.sections[0].footer.add_paragraph()._p.append(parse_xml(BOX))
    found = rules(save(d, tmp_path))
    assert {"empty-shapes", "repeating-shape"} <= found


def test_table_borders_inconsistent(tmp_path):
    d = clean_doc()
    t = d.add_table(rows=3, cols=2)  # default "Normal Table" style: no borders
    t.rows[0].cells[0].text = "x"
    rep = dr.check(save(d, tmp_path), owner=OWNER)
    hit = [f for f in rep.findings if f.rule == "table-borders"]
    assert hit and hit[0].severity == "medium"  # other tables are ruled, so it is inconsistent


def test_long_table_without_repeating_header(tmp_path):
    d = clean_doc()
    t = d.add_table(rows=9, cols=2)
    t.style = d.styles["Table Grid"]
    assert "table-header-repeat" in rules(save(d, tmp_path))


@pytest.mark.parametrize(
    "text,rule",
    [
        ("• typed bullet item", "typed-bullets"),
        ("See https://example.com/page for more.", "bare-urls"),
        ("Fill in {{CLIENT_NAME}} here.", "placeholder"),
        ("This is **bold** markdown.", "markdown-residue"),
        ("Copy with a dash — here.", "dash"),
        ("Two  spaces.", "double-space"),
    ],
)
def test_text_rules(tmp_path, text, rule):
    d = clean_doc()
    d.add_paragraph(text)
    assert rule in rules(save(d, tmp_path))


def test_blank_paragraph_runs(tmp_path):
    d = clean_doc()
    for _ in range(3):
        d.add_paragraph("")
    d.add_paragraph("after")
    assert "blank-paragraphs" in rules(save(d, tmp_path))


def test_foreign_properties_and_missing_title(tmp_path):
    d = clean_doc()
    d.core_properties.author = "Someone Else"
    d.core_properties.title = ""
    found = rules(save(d, tmp_path))
    assert {"foreign-properties", "no-title"} <= found


def test_comments_count_only_when_real(tmp_path):
    p = save(clean_doc(), tmp_path)
    with zipfile.ZipFile(p, "a") as z:  # pandoc writes an empty comments part; that is not residue
        z.writestr("word/comments.xml", '<w:comments xmlns:w="x"/>')
    assert "review-residue" not in rules(p)
    d = clean_doc()
    par = d.add_paragraph()
    par._p.append(
        parse_xml(f'<w:ins {nsdecls("w")} w:id="1" w:author="a"><w:r><w:t>x</w:t></w:r></w:ins>')
    )
    assert "review-residue" in rules(save(d, tmp_path, "tracked.docx"))


def test_unreadable_file_is_error_not_pass(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_text("not a zip")
    rep = dr.check(bad)
    assert rep.status == "error"
    assert dr.main(["check", "--gate", str(bad)]) == 1
    assert dr.main(["check", str(bad)]) == 2


def test_gate_blocks_only_on_high(tmp_path):
    d = clean_doc()
    d.add_paragraph("Two  spaces.")  # low
    assert dr.main(["check", "--gate", "--owner", OWNER, str(save(d, tmp_path))]) == 0
    d.add_paragraph("{{TODO}}")  # high
    assert dr.main(["check", "--gate", "--owner", OWNER, str(save(d, tmp_path, "h.docx"))]) == 1


def test_fix_all_clears_defects_and_keeps_words(tmp_path):
    d = clean_doc()
    sectPr = d.sections[0]._sectPr
    sectPr.remove(sectPr.find(qn("w:pgSz")))
    f = d.sections[0].footer.paragraphs[0]._p
    for r in list(f):
        if r.tag == qn("w:r"):
            f.remove(r)
    d.add_paragraph()._p.append(parse_xml(BOX))
    d.add_paragraph("Visit https://example.com today.")
    t = d.add_table(rows=9, cols=2)
    t.rows[0].cells[0].text = "head"
    src = save(d, tmp_path, "src.docx")
    out = tmp_path / "out.docx"
    res = dr.fix(
        src,
        out,
        a4=True,
        footer="",
        font="Calibri",
        borders=True,
        repeat_headers=True,
        link_urls=True,
        drop_empty_shapes=True,
    )
    assert res["words_unchanged"]
    left = rules(out)
    assert not left & {
        "page-size-missing",
        "no-page-numbers",
        "empty-shapes",
        "table-borders",
        "table-header-repeat",
        "bare-urls",
    }, left
    assert dr._plain_words(src) == dr._plain_words(out)


def test_fix_refuses_to_overwrite_input(tmp_path):
    p = save(clean_doc(), tmp_path)
    with pytest.raises(SystemExit):
        dr.fix(p, p, a4=True)


def test_fix_refuses_when_words_change(tmp_path, monkeypatch):
    p = save(clean_doc(), tmp_path)
    out = tmp_path / "out.docx"
    calls = iter([["a", "b"], ["a", "c"]])
    monkeypatch.setattr(dr, "_plain_words", lambda _p: next(calls))
    with pytest.raises(SystemExit, match="changed the document's words"):
        dr.fix(p, out, a4=True)
    assert not out.exists()


def test_render_without_renderer_is_unverified(tmp_path, monkeypatch):
    monkeypatch.setattr(dr, "_renderer", lambda: None)
    res = dr.render(save(clean_doc(), tmp_path), tmp_path / "r")
    assert res["status"] == "unverified" and res["pdf"] is None
